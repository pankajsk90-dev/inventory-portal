import sqlite3
import openpyxl
import os
import datetime

DB_PATH = 'c:/Users/Admin/Desktop/Inventory/inventory.db'
EXCEL_PATH = 'c:/Users/Admin/Downloads/INVENTORY ALL PRODUCTS ALL PLATFORMS - Category mapping fixed (1).xlsx'

def init_db():
    if os.path.exists(DB_PATH):
        try:
            os.remove(DB_PATH)
        except:
            pass
    
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    
    # Products table
    cur.execute('''
    CREATE TABLE IF NOT EXISTS products (
        sku TEXT PRIMARY KEY,
        name TEXT NOT NULL,
        category TEXT,
        brand TEXT NOT NULL
    )
    ''')
    
    # Inventory Stock & Thresholds
    cur.execute('''
    CREATE TABLE IF NOT EXISTS inventory_stock (
        sku TEXT PRIMARY KEY,
        total_inventory INTEGER NOT NULL DEFAULT 0,
        reorder_threshold INTEGER NOT NULL DEFAULT 0,
        FOREIGN KEY (sku) REFERENCES products(sku)
    )
    ''')
    
    # Dispatch Online
    cur.execute('''
    CREATE TABLE IF NOT EXISTS dispatch_online (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        dispatch_date TEXT,
        brand TEXT,
        product_name TEXT,
        sku TEXT,
        category TEXT,
        platform TEXT,
        location TEXT,
        po_number TEXT,
        po_quantity INTEGER DEFAULT 0,
        actual_sent INTEGER NOT NULL DEFAULT 0,
        appointment_date TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    )
    ''')
    
    # Dispatch GT MT
    cur.execute('''
    CREATE TABLE IF NOT EXISTS dispatch_gt_mt (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        dispatch_date TEXT,
        brand TEXT,
        product_name TEXT,
        sku TEXT,
        category TEXT,
        channel TEXT,
        buyer_distributor TEXT,
        location TEXT,
        po_number TEXT,
        invoice_no TEXT,
        po_quantity INTEGER DEFAULT 0,
        actual_sent INTEGER NOT NULL DEFAULT 0,
        appointment_date TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    )
    ''')
    
    # Adjustments
    cur.execute('''
    CREATE TABLE IF NOT EXISTS adjustments (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        adjustment_date TEXT,
        brand TEXT,
        channel TEXT,
        product_name TEXT,
        sku TEXT,
        category TEXT,
        adjustment_type TEXT,
        quantity INTEGER NOT NULL DEFAULT 0,
        reason TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    )
    ''')
    
    # Audit log
    cur.execute('''
    CREATE TABLE IF NOT EXISTS audit_log (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        timestamp TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        action TEXT,
        details TEXT
    )
    ''')
    
    conn.commit()
    return conn

def format_date(val):
    if val is None:
        return ''
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime('%Y-%m-%d')
    s = str(val).strip()
    if s:
        for fmt in ('%Y-%m-%d %H:%M:%S', '%d/%m/%y', '%d/%m/%Y', '%Y-%m-%d'):
            try:
                dt = datetime.datetime.strptime(s, fmt)
                return dt.strftime('%Y-%m-%d')
            except ValueError:
                pass
    return s

def ingest_excel(conn):
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    cur = conn.cursor()
    
    # 1. Product Master
    ws_prod = wb['Product Master']
    prod_map = {}
    sku_map = {}
    
    for r in range(5, ws_prod.max_row + 1):
        name = ws_prod.cell(r, 1).value
        sku = ws_prod.cell(r, 2).value
        cat = ws_prod.cell(r, 3).value
        brand = ws_prod.cell(r, 4).value
        if sku and str(sku).strip():
            sku = str(sku).strip()
            name = str(name).strip() if name else ''
            cat = str(cat).strip() if cat else ''
            brand = str(brand).strip() if brand else ''
            
            cur.execute('INSERT OR REPLACE INTO products (sku, name, category, brand) VALUES (?, ?, ?, ?)',
                        (sku, name, cat, brand))
            prod_map[name.lower()] = {'sku': sku, 'name': name, 'category': cat, 'brand': brand}
            sku_map[sku] = {'name': name, 'category': cat, 'brand': brand}
    
    print(f'Ingested {len(sku_map)} products.')
    
    # 2. Inventory Stock & Thresholds
    ws_stock = wb['Inventory Stock']
    stock_count = 0
    for r in range(5, ws_stock.max_row + 1):
        sku = ws_stock.cell(r, 1).value
        if sku and str(sku).strip() and str(sku).strip() != 'None' and not str(sku).startswith('GRAND'):
            sku = str(sku).strip()
            tot_inv = ws_stock.cell(r, 5).value or 0
            reorder = ws_stock.cell(r, 6).value or 0
            try:
                tot_inv = int(tot_inv)
            except:
                tot_inv = 0
            try:
                reorder = int(reorder)
            except:
                reorder = 0
            cur.execute('INSERT OR REPLACE INTO inventory_stock (sku, total_inventory, reorder_threshold) VALUES (?, ?, ?)',
                        (sku, tot_inv, reorder))
            stock_count += 1
    print(f'Ingested {stock_count} stock thresholds.')
    
    # 3. Dispatch online
    ws_disp = wb['Dispatch online']
    disp_online_count = 0
    for r in range(5, ws_disp.max_row + 1):
        d_val = ws_disp.cell(r, 1).value
        brand = ws_disp.cell(r, 2).value
        pname = ws_disp.cell(r, 3).value
        plat = ws_disp.cell(r, 4).value
        loc = ws_disp.cell(r, 5).value
        units = ws_disp.cell(r, 6).value
        po_no = ws_disp.cell(r, 7).value
        appt = ws_disp.cell(r, 8).value
        sku = ws_disp.cell(r, 9).value
        cat = ws_disp.cell(r, 10).value
        
        if not pname and not units and not po_no:
            continue
            
        pname_clean = str(pname).strip() if pname else ''
        sku_clean = str(sku).strip() if sku else ''
        
        if not sku_clean and pname_clean.lower() in prod_map:
            sku_clean = prod_map[pname_clean.lower()]['sku']
        if not brand and sku_clean in sku_map:
            brand = sku_map[sku_clean]['brand']
        if not cat and sku_clean in sku_map:
            cat = sku_map[sku_clean]['category']
            
        try:
            actual_sent = int(units or 0)
        except:
            actual_sent = 0
            
        po_qty = actual_sent
        po_num_str = str(po_no).strip() if po_no is not None else ''
        
        cur.execute('''
        INSERT INTO dispatch_online (
            dispatch_date, brand, product_name, sku, category, platform, location, po_number, po_quantity, actual_sent, appointment_date
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            format_date(d_val),
            str(brand).strip() if brand else '',
            pname_clean,
            sku_clean,
            str(cat).strip() if cat else '',
            str(plat).strip() if plat else '',
            str(loc).strip() if loc else '',
            po_num_str,
            po_qty,
            actual_sent,
            format_date(appt)
        ))
        disp_online_count += 1
    print(f'Ingested {disp_online_count} online dispatches.')
    
    # 4. Dispatch GT MT
    ws_gtmt = wb['Dispatch GT MT']
    gtmt_count = 0
    for r in range(5, ws_gtmt.max_row + 1):
        d_val = ws_gtmt.cell(r, 1).value
        brand = ws_gtmt.cell(r, 2).value
        pname = ws_gtmt.cell(r, 3).value
        channel = ws_gtmt.cell(r, 4).value
        buyer = ws_gtmt.cell(r, 5).value
        loc = ws_gtmt.cell(r, 6).value
        units = ws_gtmt.cell(r, 7).value
        inv_no = ws_gtmt.cell(r, 8).value
        appt = ws_gtmt.cell(r, 9).value
        sku = ws_gtmt.cell(r, 10).value
        cat = ws_gtmt.cell(r, 11).value
        
        if not pname and not units and not buyer:
            continue
            
        pname_clean = str(pname).strip() if pname else ''
        sku_clean = str(sku).strip() if sku else ''
        if not sku_clean and pname_clean.lower() in prod_map:
            sku_clean = prod_map[pname_clean.lower()]['sku']
        if not brand and sku_clean in sku_map:
            brand = sku_map[sku_clean]['brand']
        if not cat and sku_clean in sku_map:
            cat = sku_map[sku_clean]['category']
            
        try:
            actual_sent = int(units or 0)
        except:
            actual_sent = 0
            
        cur.execute('''
        INSERT INTO dispatch_gt_mt (
            dispatch_date, brand, product_name, sku, category, channel, buyer_distributor, location, po_number, invoice_no, po_quantity, actual_sent, appointment_date
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            format_date(d_val),
            str(brand).strip() if brand else '',
            pname_clean,
            sku_clean,
            str(cat).strip() if cat else '',
            str(channel).strip() if channel else 'GT',
            str(buyer).strip() if buyer else '',
            str(loc).strip() if loc else '',
            '',
            str(inv_no).strip() if inv_no is not None else '',
            actual_sent,
            actual_sent,
            format_date(appt)
        ))
        gtmt_count += 1
    print(f'Ingested {gtmt_count} GT/MT dispatches.')
    
    # 5. Adjustments
    ws_adj = wb['Adjustments']
    adj_count = 0
    for r in range(5, ws_adj.max_row + 1):
        d_val = ws_adj.cell(r, 1).value
        brand = ws_adj.cell(r, 2).value
        chan = ws_adj.cell(r, 3).value
        pname = ws_adj.cell(r, 4).value
        atype = ws_adj.cell(r, 5).value
        qty = ws_adj.cell(r, 6).value
        reason = ws_adj.cell(r, 7).value
        sku = ws_adj.cell(r, 8).value
        cat = ws_adj.cell(r, 9).value
        
        if not pname and not qty and not atype:
            continue
            
        pname_clean = str(pname).strip() if pname else ''
        sku_clean = str(sku).strip() if sku else ''
        if not sku_clean and pname_clean.lower() in prod_map:
            sku_clean = prod_map[pname_clean.lower()]['sku']
        if not brand and sku_clean in sku_map:
            brand = sku_map[sku_clean]['brand']
        if not cat and sku_clean in sku_map:
            cat = sku_map[sku_clean]['category']
            
        try:
            quantity = int(qty or 0)
        except:
            quantity = 0
            
        cur.execute('''
        INSERT INTO adjustments (
            adjustment_date, brand, channel, product_name, sku, category, adjustment_type, quantity, reason
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            format_date(d_val),
            str(brand).strip() if brand else '',
            str(chan).strip() if chan else '',
            pname_clean,
            sku_clean,
            str(cat).strip() if cat else '',
            str(atype).strip() if atype else 'Other',
            quantity,
            str(reason).strip() if reason else ''
        ))
        adj_count += 1
    print(f'Ingested {adj_count} adjustments.')
    
    cur.execute("INSERT INTO audit_log (action, details) VALUES ('INITIAL_INGESTION', 'Imported initial data from Excel')")
    conn.commit()
    wb.close()
    print('All data ingested successfully.')

if __name__ == '__main__':
    conn = init_db()
    ingest_excel(conn)
    conn.close()
